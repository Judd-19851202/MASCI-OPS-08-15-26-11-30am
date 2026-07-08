"""Track 24.13 live smoke tests against the real REACT_APP_BACKEND_URL ingress.

Covers:
- GET /api/dr-v2/meta health smoke
- GET /api/daily-reports/{id}/evidence-manifest structure (v24.13.1 + sha256 hash prefix)
- POST /api/daily-reports/evidence/extract on real TXT, PDF (text), scanned PDF,
  corrupt PDF, unsupported .doc, real XLSX, real DOCX, real CSV
- Full DR POST with evidence_manifest and PDF section 10A + 10B verification
- Legacy DR without evidence_manifest PDF has no 10B
"""
import base64
import io
import os
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    # fallback for direct test runs
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                BASE_URL = line.split("=", 1)[1].strip().rstrip("/")

ADMIN_EMAIL = "jaymn.judd@mascigc.com"
ADMIN_PASSWORD = "Maddix123!"


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(
        f"{BASE_URL}/api/auth/multi-login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        timeout=30,
    )
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    data = r.json()
    tok = (data.get("portal_tokens") or {}).get("admin")
    assert tok, f"no admin token: {data}"
    return tok


@pytest.fixture(scope="module")
def any_dr_id(admin_token):
    r = requests.get(
        f"{BASE_URL}/api/daily-reports?limit=5",
        headers={"X-Admin-Token": admin_token},
        timeout=30,
    )
    assert r.status_code == 200, r.text
    items = r.json()
    if isinstance(items, dict):
        items = items.get("items") or items.get("daily_reports") or []
    assert items, "no DRs available"
    return items[0].get("id") or items[0].get("report_id")


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


def test_meta_endpoint():
    last = None
    for _ in range(3):
        try:
            r = requests.get(f"{BASE_URL}/api/dr-v2/meta", timeout=60)
            last = r
            if r.status_code == 200:
                break
        except requests.RequestException as e:
            last = e
    assert getattr(last, "status_code", None) == 200, f"meta failed: {last}"
    body = last.json()
    assert "agents" in body


def test_evidence_manifest_endpoint(admin_token, any_dr_id):
    r = requests.get(
        f"{BASE_URL}/api/daily-reports/{any_dr_id}/evidence-manifest",
        headers={"X-Admin-Token": admin_token},
        timeout=30,
    )
    assert r.status_code == 200, r.text
    m = r.json()
    assert m.get("version") == "24.13.1", f"version={m.get('version')}"
    h = m.get("manifest_hash", "")
    assert h.startswith("sha256:"), f"hash={h}"
    for key in ("photos", "attachments", "typed_fields", "material_reconciliation", "warnings"):
        assert key in m, f"missing key: {key}"


def _extract(payload):
    # translate our normalized keys to the endpoint's payload shape
    body = {
        "filename": payload.get("filename", ""),
        "mime": payload.get("content_type", ""),
        "data_base64": payload.get("content_base64", ""),
    }
    return requests.post(
        f"{BASE_URL}/api/daily-reports/evidence/extract",
        json=body,
        timeout=60,
    )


def test_extract_txt():
    r = _extract({
        "filename": "note.txt",
        "content_type": "text/plain",
        "content_base64": _b64(b"Hello Track 24.13 evidence."),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extracted"
    assert "Track 24.13" in (d.get("text") or "")


def test_extract_pdf_with_text():
    try:
        import fitz  # PyMuPDF
    except ImportError:
        pytest.skip("pymupdf not available")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Track2413PDFMarker payload text")
    buf = doc.tobytes()
    doc.close()
    r = _extract({
        "filename": "test.pdf",
        "content_type": "application/pdf",
        "content_base64": _b64(buf),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extracted", d
    assert d.get("page_count", 0) >= 1
    assert "Track2413PDFMarker" in (d.get("text") or "")


def test_extract_pdf_scanned_no_text():
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not available")
    doc = fitz.open()
    doc.new_page()  # blank page, no text layer
    buf = doc.tobytes()
    doc.close()
    r = _extract({
        "filename": "scan.pdf",
        "content_type": "application/pdf",
        "content_base64": _b64(buf),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "scanned_pdf_no_text", d


def test_extract_corrupt_pdf():
    r = _extract({
        "filename": "corrupt.pdf",
        "content_type": "application/pdf",
        "content_base64": _b64(b"%PDF-1.4 not-a-real-pdf-content-garbage"),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "corrupt", d


def test_extract_doc_unsupported():
    r = _extract({
        "filename": "old.doc",
        "content_type": "application/msword",
        "content_base64": _b64(b"\xd0\xcf\x11\xe0legacy-doc-bytes"),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "unsupported", d
    assert d.get("reason")


def test_extract_xlsx():
    try:
        from openpyxl import Workbook
    except ImportError:
        pytest.skip("openpyxl not available")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.append(["ticket_no", "supplier", "material", "qty", "unit"])
    ws.append(["T-001", "AcmeAgg", "gravel", 12.5, "ton"])
    ws.append(["T-002", "AcmeAgg", "sand", 8.0, "ton"])
    buf = io.BytesIO()
    wb.save(buf)
    r = _extract({
        "filename": "book.xlsx",
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "content_base64": _b64(buf.getvalue()),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extracted", d
    assert d.get("sheet_names")
    rows = d.get("rows") or []
    assert rows, "rows empty"
    assert any("[[SHEET:" in (r0 if isinstance(r0, str) else " ".join(map(str, r0))) for r0 in rows), rows[:3]


def test_extract_docx():
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not available")
    doc = Document()
    doc.add_paragraph("Track2413DocxMarker body content.")
    buf = io.BytesIO()
    doc.save(buf)
    r = _extract({
        "filename": "note.docx",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "content_base64": _b64(buf.getvalue()),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extracted", d
    assert "Track2413DocxMarker" in (d.get("text") or "")


def test_extract_csv():
    csv_body = b"ticket_no,supplier,material,qty,unit\nT-100,AcmeAgg,gravel,10,ton\nT-101,AcmeAgg,sand,5,ton\n"
    r = _extract({
        "filename": "tickets.csv",
        "content_type": "text/csv",
        "content_base64": _b64(csv_body),
    })
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "extracted", d
    row_count = d.get("row_count", 0) or (d.get("ext_meta") or {}).get("rows_seen", 0) or len(d.get("rows") or [])
    assert row_count >= 3, d
    # headers preserved somewhere
    rows = d.get("rows") or []
    text_blob = (d.get("text") or "") + " " + " ".join(str(x) for x in rows)
    assert "ticket_no" in text_blob


def test_dr_with_evidence_manifest_and_pdf_10a_10b(admin_token):
    payload = {
        "project_name": "Smoke Track 24.13",
        "project_code": "SMOKE-24-13",
        "location": "Smoke Test Site",
        "report_date": "2026-01-15",
        "date": "2026-01-15",
        "prepared_by": "Smoke Tester",
        "supervisor_name": "Smoke Tester",
        "crew_notes": "smoke test crew",
        "ai_accepted_summary": "Test accepted operational summary for 10A hero.",
        "ai_accepted_summary_meta": {"agent": "manifest_summary", "source": "supervisor_accepted"},
        "evidence_manifest": {
            "version": "24.13.1",
            "manifest_hash": "sha256:smoketest",
            "photos": [],
            "attachments": [
                {
                    "filename": "tickets.csv",
                    "content_type": "text/csv",
                    "size": 100,
                    "extraction_status": "extracted",
                    "row_count": 3,
                }
            ],
            "typed_fields": {},
            "material_reconciliation": {
                "matched": [{"ticket_no": "T-100", "material": "gravel", "qty": 10, "unit": "ton"}],
                "variance_flags": [],
                "summary": "1 ticket matched",
            },
            "warnings": ["Evidence smoke warning"],
        },
    }
    r = requests.post(
        f"{BASE_URL}/api/daily-reports",
        json=payload,
        headers={"X-Admin-Token": admin_token},
        timeout=60,
    )
    assert r.status_code in (200, 201), r.text
    dr = r.json()
    rid = dr.get("id") or dr.get("report_id")
    assert rid

    # fetch pdf
    pdf_r = requests.get(
        f"{BASE_URL}/api/daily-reports/{rid}/pdf",
        headers={"X-Admin-Token": admin_token},
        timeout=60,
    )
    assert pdf_r.status_code == 200, pdf_r.text[:500]
    assert pdf_r.content[:4] == b"%PDF", pdf_r.content[:10]

    # extract text
    try:
        import fitz
        doc = fitz.open(stream=pdf_r.content, filetype="pdf")
        text = "\n".join(p.get_text() for p in doc)
        doc.close()
    except ImportError:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_r.content))
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages)

    assert "OPERATIONAL INTELLIGENCE SUMMARY" in text, text[:2000]
    assert "ATTACHMENT" in text and "EVIDENCE" in text, "10B missing"
    assert "tickets.csv" in text or "T-100" in text, "attachment row / ticket missing"
    assert "Evidence smoke warning" in text or "smoke warning" in text.lower(), "warning missing"


def test_legacy_dr_pdf_no_10b(admin_token):
    payload = {
        "project_name": "Smoke Legacy",
        "project_code": "SMOKE-24-13-LEGACY",
        "location": "Legacy Site",
        "report_date": "2026-01-15",
        "date": "2026-01-15",
        "prepared_by": "Legacy Smoke",
        "supervisor_name": "Legacy Smoke",
        "crew_notes": "legacy",
    }
    r = requests.post(
        f"{BASE_URL}/api/daily-reports",
        json=payload,
        headers={"X-Admin-Token": admin_token},
        timeout=60,
    )
    assert r.status_code in (200, 201), r.text
    rid = r.json().get("id") or r.json().get("report_id")
    pdf_r = requests.get(
        f"{BASE_URL}/api/daily-reports/{rid}/pdf",
        headers={"X-Admin-Token": admin_token},
        timeout=60,
    )
    assert pdf_r.status_code == 200
    try:
        import fitz
        doc = fitz.open(stream=pdf_r.content, filetype="pdf")
        text = "\n".join(p.get_text() for p in doc)
        doc.close()
    except ImportError:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_r.content))
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
    assert "10B" not in text and "ATTACHMENT & DOCUMENT EVIDENCE" not in text, "10B leaked into legacy DR"
