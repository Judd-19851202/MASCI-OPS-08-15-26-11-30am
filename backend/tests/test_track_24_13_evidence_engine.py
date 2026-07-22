"""TRACK 24.13 · Daily Report Evidence Intelligence Engine · locks.

Static-and-offline test suite for the Evidence Intelligence Engine.
No live LLM calls · no live R2 · no live Mongo mutations. Every test
is executable in the CI matrix in under one second.
"""
from __future__ import annotations

import io
from pathlib import Path

# ── Test 1 · Extraction dispatcher covers every advertised type ───

def test_extraction_status_vocabulary_locked():
    from services.dr_evidence.extract import EXTRACTION_STATUSES
    expected = {
        "not_started", "extracted", "unsupported", "failed",
        "too_large", "encrypted", "corrupt", "scanned_pdf_no_text",
    }
    assert set(EXTRACTION_STATUSES) == expected, (
        f"TRACK 24.13 · extraction status vocab drifted: got={EXTRACTION_STATUSES}"
    )


def test_extract_txt_ok():
    from services.dr_evidence.extract import extract_attachment
    r = extract_attachment(
        filename="notes.txt", mime="text/plain",
        data=b"Line one\nLine two",
    )
    assert r.status == "extracted"
    assert "Line one" in r.text
    assert r.confidence >= 0.8


def test_extract_csv_ok_and_headers_survive():
    from services.dr_evidence.extract import extract_attachment
    r = extract_attachment(
        filename="tix.csv", mime="text/csv",
        data=b"ticket,supplier,material,qty,unit\n"
             b"1001,ACME,Base Rock,25.5,tons\n",
    )
    assert r.status == "extracted"
    assert r.rows
    assert r.rows[0][0].lower().startswith("ticket")


def test_extract_pdf_real_text_extracted():
    import fitz  # PyMuPDF
    from services.dr_evidence.extract import extract_attachment
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 80), "Permit #P-42\nProject 24-13")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    doc.close()
    r = extract_attachment(
        filename="permit.pdf", mime="application/pdf", data=data,
    )
    assert r.status == "extracted"
    assert "Permit" in r.text
    assert r.page_count == 1


def test_extract_pdf_scanned_no_text():
    import fitz
    from services.dr_evidence.extract import extract_attachment
    doc = fitz.open()
    doc.new_page()  # no text at all
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    doc.close()
    r = extract_attachment(
        filename="scan.pdf", mime="application/pdf", data=data,
    )
    assert r.status == "scanned_pdf_no_text", (
        "PDFs with no embedded text must be flagged so the AI does "
        "not hallucinate their contents."
    )


def test_extract_pdf_corrupt_flagged():
    from services.dr_evidence.extract import extract_attachment
    r = extract_attachment(
        filename="bad.pdf", mime="application/pdf",
        data=b"not really a pdf",
    )
    assert r.status == "corrupt"


def test_extract_docx_ok():
    import docx
    from services.dr_evidence.extract import extract_attachment
    d = docx.Document()
    d.add_paragraph("Inspection passed at STA 100+00")
    buf = io.BytesIO()
    d.save(buf)
    r = extract_attachment(filename="note.docx", mime="", data=buf.getvalue())
    assert r.status == "extracted"
    assert "Inspection" in r.text


def test_extract_xlsx_ok_and_carries_sheet_marker():
    import openpyxl
    from services.dr_evidence.extract import extract_attachment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.append(["ticket_no", "material", "qty", "unit"])
    ws.append([501, "Base Rock", 20.0, "tons"])
    buf = io.BytesIO()
    wb.save(buf)
    r = extract_attachment(filename="t.xlsx", mime="", data=buf.getvalue())
    assert r.status == "extracted"
    assert r.sheet_names == ["Tickets"]
    assert r.rows[0][0].startswith("[[SHEET:")


def test_extract_legacy_doc_marked_unsupported():
    from services.dr_evidence.extract import extract_attachment
    r = extract_attachment(
        filename="thing.doc", mime="application/msword",
        data=b"\x00\x01\x02",
    )
    assert r.status == "unsupported"


def test_extract_too_large_is_bounded():
    from services.dr_evidence.extract import extract_attachment, MAX_BYTES
    r = extract_attachment(
        filename="huge.pdf", mime="application/pdf",
        data=b"x" * (MAX_BYTES + 1),
    )
    assert r.status == "too_large"


def test_hash_bytes_stable_and_prefixed():
    from services.dr_evidence.extract import hash_bytes
    h1 = hash_bytes(b"hello")
    h2 = hash_bytes(b"hello")
    assert h1 == h2 and h1.startswith("sha256:")


# ── Materials reconciliation ────────────────────────────────────────

def test_normalize_ticket_row_maps_common_headers():
    from services.dr_evidence.materials import normalize_ticket_row
    t = normalize_ticket_row(
        {"Ticket #": "9001", "Vendor": "ACME", "Product": "Asphalt",
         "Tons": "12.5", "UOM": "tons", "Truck": "T-42"},
        source="csv",
    )
    assert t.ticket_number == "9001"
    assert t.supplier == "ACME"
    assert t.material == "Asphalt"
    assert t.quantity == 12.5
    assert t.unit == "tons"
    assert t.truck == "T-42"


def test_reconcile_tickets_matches_and_flags_variance():
    from services.dr_evidence.materials import (
        normalize_ticket_row, tickets_from_rows, reconcile_tickets,
    )
    entered = [normalize_ticket_row(
        {"ticket_number": "1001", "material": "Base Rock",
         "quantity": 25.5, "unit": "tons"}, source="entered",
    )]
    rows = [
        ["ticket_number", "supplier", "material", "quantity", "unit"],
        ["1001", "ACME", "Base Rock", "25.5", "tons"],
        ["9999", "ACME", "Base Rock", "60.0", "tons"],  # unmatched + variance
    ]
    extracted = tickets_from_rows(rows, source="csv")
    result = reconcile_tickets(entered, extracted)
    assert len(result.matched) == 1
    assert len(result.unmatched_extracted) == 1
    assert any("variance" in a for a in result.advisories)


# ── Manifest builder ────────────────────────────────────────────────

def test_manifest_includes_typed_fields_and_attachments():
    from services.dr_evidence import build_manifest, manifest_to_ai_bundle
    from services.dr_evidence.extract import ExtractionResult
    report = {
        "doc_id": "DR-2026-99999",
        "project_number": "24-13",
        "project_name": "T24-13 CERT",
        "client": "MASCI",
        "project_manager": "PM Test",
        "report_date": "2026-02-15",
        "supervisor_name": "Foreman F",
        "masci_crews": [{"trade": "Excavation", "count": 6}],
        "equipment": [{"unit": "EX-01", "hours": 8}],
        "materials": [
            {"ticket_number": "1001", "material": "Base Rock",
             "quantity": 25.5, "unit": "tons"},
        ],
        "excavation": {"depth_ft": 6, "safe_to_use": True},
        "photos": ["photo://x/y.jpg"],
    }
    att_ext = [{
        "id": "att-1", "filename": "tickets.csv",
        "mime": "text/csv", "size_bytes": 500, "source_section": "materials",
        "extraction": ExtractionResult(
            status="extracted", confidence=0.85,
            rows=[
                ["ticket_number", "material", "quantity", "unit"],
                ["1001", "Base Rock", "25.5", "tons"],
                ["8888", "Base Rock", "40", "tons"],
            ],
            text="ticket_number\tmaterial…",
        ).to_dict(),
    }]
    m = build_manifest(report, attachment_extractions=att_ext)
    d = m.to_dict()
    assert d["project_number"] == "24-13"
    assert d["typed_fields"].get("masci_crews")
    assert d["typed_fields"].get("excavation")
    assert d["attachments"][0]["extraction_status"] == "extracted"
    # Photos are present but analysis_status may be "not_started" — no
    # intel provided in this test — which is honest.
    assert d["photos"][0]["analysis_status"] in ("not_started", "unavailable")
    # Reconciliation produced advisory rows.
    recon = d["material_reconciliation"]
    assert recon["matched"], "matched ticket 1001 must appear"
    bundle = manifest_to_ai_bundle(m)
    assert "typed_fields" in bundle and "attachments" in bundle
    assert bundle["material_reconciliation"] is recon


def test_manifest_hash_is_stable_across_timestamps():
    from services.dr_evidence import build_manifest, manifest_hash
    report = {"doc_id": "DR-X", "project_number": "P", "materials": []}
    h1 = manifest_hash(build_manifest(report))
    h2 = manifest_hash(build_manifest(report))
    assert h1 == h2, (
        "TRACK 24.13 · manifest_hash must exclude timestamps + "
        "confidence floats so the AI cache is not busted spuriously."
    )


def test_manifest_flags_warnings_for_unextracted_attachments():
    from services.dr_evidence import build_manifest
    from services.dr_evidence.extract import ExtractionResult
    m = build_manifest(
        {"doc_id": "DR-Y"},
        attachment_extractions=[{
            "id": "a1", "filename": "scan.pdf",
            "extraction": ExtractionResult(
                status="scanned_pdf_no_text", reason="no_embedded_text",
            ).to_dict(),
        }],
    )
    assert any("scan.pdf" in w for w in m.warnings)


# ── AI prompt hardening ────────────────────────────────────────────

def test_manifest_summary_prompt_has_evidence_class_rules():
    from services.dr_ai.agents import AGENTS
    assert "manifest_summary" in AGENTS, (
        "TRACK 24.13 · `manifest_summary` agent must be registered."
    )
    sys_prompt = AGENTS["manifest_summary"]["system"]
    # Anti-hallucination requirements enforced by the prompt itself.
    for phrase in [
        "extraction_status",
        "analysis_status",
        "material_reconciliation",
        "STRICT JSON",
        "photo_and_attachment_evidence",
        "materials_and_tickets",
        "safety_and_quality",
        "warnings",
        "evidence_refs",
    ]:
        assert phrase in sys_prompt, (
            f"TRACK 24.13 · manifest_summary prompt missing required "
            f"phrase: {phrase!r}"
        )
    # Explicit anti-guessing.
    lower = sys_prompt.lower()
    assert "do not guess file contents" in lower
    assert "hallucinat" not in lower or "do not" in lower


# ── PDF attachment evidence section ────────────────────────────────

def test_pdf_attachment_evidence_section_hidden_when_no_manifest():
    from pdf_render import _render_attachment_evidence_section
    assert _render_attachment_evidence_section({}) == ""


def test_pdf_attachment_evidence_section_renders_when_manifest_present():
    from pdf_render import _render_attachment_evidence_section
    d = {
        "evidence_manifest": {
            "attachments": [
                {"filename": "permit.pdf", "extraction_status": "extracted",
                 "page_count": 2, "row_count": 0, "source_section": "safety"},
                {"filename": "scan.pdf",
                 "extraction_status": "scanned_pdf_no_text",
                 "extraction_reason": "no_embedded_text"},
            ],
            "material_reconciliation": {
                "matched": [{"ticket_number": "1001"}],
                "unmatched_extracted": [],
                "advisories": ["Quantity delta on base rock: entered 25 vs extracted 40."],
            },
            "warnings": ["scan.pdf: scanned_pdf_no_text"],
        },
    }
    html = _render_attachment_evidence_section(d)
    assert "permit.pdf" in html
    assert "scanned_pdf_no_text" in html
    assert "Advisories" in html or "advisory" in html.lower()
    assert "Quantity delta" in html
    assert "Evidence Warnings" in html


def test_pdf_attachment_evidence_section_falls_back_to_raw_attachment_refs():
    from pdf_render import _render_attachment_evidence_section

    d = {
        "attachments": [
            {
                "filename": "daily_ticket.pdf",
                "category": "PDF",
                "uploaded_at": "2026-07-22T23:14:30Z",
                "attachment_ref": "photo://masci-hub/documents/2026/07/dr_attachment_daily_ticket.pdf",
            },
            {
                "filename": "notes with spaces & punctuation!.txt",
                "category": "Document",
                "uploaded_at": "2026-07-22T23:14:31Z",
            },
        ],
    }

    html = _render_attachment_evidence_section(d)
    assert "daily_ticket.pdf" in html
    assert "notes with spaces &amp; punctuation!.txt" in html
    assert "Attachment status" in html
    assert "Saved · photo://masci-hub/documents/2026/07/dr_attach" in html


# ── V1 DR shape carries `evidence_manifest` field ──────────────────

def test_daily_report_model_accepts_evidence_manifest():
    from routes.daily_reports import DailyReportCreate
    dr = DailyReportCreate(
        project_name="T", location="L", report_date="2026-02-15",
        prepared_by="F", evidence_manifest={"version": "24.13.1"},
    )
    d = dr.model_dump()
    assert d["evidence_manifest"] == {"version": "24.13.1"}


# ── Static: no secrets/prompt leaks in the manifest routes ─────────

def test_no_provider_key_leak_in_dr_evidence_source():
    root = Path("/app/backend/services/dr_evidence")
    banned = ("EMERGENT_LLM_KEY", "ANTHROPIC_API_KEY", "OPENAI_API_KEY",
              "gemini_api_key")
    for p in root.rglob("*.py"):
        src = p.read_text(encoding="utf-8")
        for b in banned:
            assert b not in src, (
                f"TRACK 24.13 · secret name {b!r} appears in {p} — "
                "evidence engine must be provider-agnostic."
            )
